from io import BytesIO

from django.contrib.auth.models import Permission, User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from django.core.files.uploadedfile import SimpleUploadedFile

from .models import Import, ImportItem, ImportType
from .services import process_xlsx_import
from .tests import xlsx_file


class ImportDownloadTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('download-reader')
        cls.user.user_permissions.add(Permission.objects.get(codename='view_import'))
        cls.columns = ['ДБЗ', 'ИИН', 'ФИО', 'Сумма']
        cls.record = Import.objects.create(import_type=ImportType.objects.get(code='contracts'),
                                          file_name='Договоры.xlsx', metadata={'columns': cls.columns})
        cls.first = ImportItem.objects.create(import_record=cls.record, row_number=2,
                                             data={'ДБЗ': '001', 'ИИН': '000000000001', 'ФИО': '=1+1', 'Сумма': '123.45'},
                                             status=ImportItem.Status.PROCESSED)
        cls.failed = ImportItem.objects.create(import_record=cls.record, row_number=4,
                                              data={'ДБЗ': '002', 'ИИН': 'неверный', 'ФИО': 'Тест'},
                                              status=ImportItem.Status.FAILED)

    def setUp(self):
        self.client.force_login(self.user)

    def download(self):
        response = self.client.get(reverse('imports:download', args=[self.record.pk]))
        self.assertEqual(response.status_code, 200)
        content = b''.join(response.streaming_content)
        response.close()
        return response, load_workbook(BytesIO(content))

    def test_saved_rows_headers_and_identifiers_are_preserved(self):
        response, book = self.download()
        self.assertIn('attachment;', response['Content-Disposition'])
        self.assertIn('filename*=', response['Content-Disposition'])
        self.assertEqual(response['Cache-Control'], 'no-store')
        sheet = book.active
        self.assertEqual([cell.value for cell in sheet[1]], self.columns)
        self.assertEqual(sheet['A2'].value, '001')
        self.assertEqual(sheet['B2'].value, '000000000001')
        self.assertEqual(sheet['C2'].value, '=1+1')
        self.assertEqual(sheet['C2'].data_type, 's')
        self.assertEqual(sheet['D2'].value, '123.45')
        self.assertIsNone(sheet['A3'].value)
        self.assertEqual(sheet['B4'].value, 'неверный')
        self.assertIsNone(sheet['D4'].value)
        book.close()

    def test_each_download_uses_current_saved_data(self):
        _, first = self.download()
        self.first.data['ФИО'] = 'Обновлённое имя'
        self.first.save(update_fields=['data'])
        _, second = self.download()
        self.assertEqual(first.active['C2'].value, '=1+1')
        self.assertEqual(second.active['C2'].value, 'Обновлённое имя')
        first.close()
        second.close()

    def test_legacy_import_without_metadata_uses_saved_headers(self):
        self.record.metadata = {}
        self.record.save(update_fields=['metadata'])
        _, book = self.download()
        self.assertEqual([cell.value for cell in book.active[1]], self.columns)
        book.close()

    def test_empty_import_downloads_header_only_workbook(self):
        self.record.items.all().delete()
        _, book = self.download()
        self.assertEqual(book.active.max_row, 1)
        self.assertEqual([cell.value for cell in book.active[1]], self.columns)
        book.close()

    def test_download_requires_view_permission_and_existing_import(self):
        self.assertEqual(self.client.get(reverse('imports:download', args=[999999])).status_code, 404)
        self.assertEqual(self.client.post(reverse('imports:download', args=[self.record.pk])).status_code, 405)
        self.client.force_login(User.objects.create_user('no-download-access'))
        self.assertEqual(self.client.get(reverse('imports:download', args=[self.record.pk])).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('imports:download', args=[self.record.pk])).status_code, 302)

    def test_generated_workbook_can_be_imported_again(self):
        self.record = Import.objects.create(import_type=self.record.import_type)
        process_xlsx_import(self.record, xlsx_file(
            ['ДБЗ', 'ИИН', 'ФИО', 'Основной долг (выкуп)'],
            [['ROUNDTRIP-01', '000000000061', 'Тестовый заёмщик', '123.45']],
        ), preview_only=True)
        self.assertEqual(self.record.successful_items, 1)
        response = self.client.get(reverse('imports:download', args=[self.record.pk]))
        content = b''.join(response.streaming_content)
        response.close()
        regenerated = Import.objects.create(import_type=self.record.import_type)
        process_xlsx_import(regenerated, SimpleUploadedFile('regenerated.xlsx', content), preview_only=True)
        self.assertEqual(regenerated.successful_items, 1)
        self.assertEqual(regenerated.failed_items, 0)
        self.assertEqual(regenerated.items.get().data, self.record.items.get().data)

    def test_all_column_widths_fit_headers_and_content_in_any_row(self):
        columns = ['Очень длинное название колонки', 'Комментарий'] + [f'Колонка {n}' for n in range(3, 36)]
        self.record.metadata = {'columns': columns}
        self.record.save(update_fields=['metadata'])
        self.first.data = {column: '1' for column in columns}
        self.first.save(update_fields=['data'])
        long_text = 'Длинный комментарий из последней строки импорта'
        self.failed.data = {columns[1]: long_text, columns[-1]: long_text * 2}
        self.failed.save(update_fields=['data'])
        _, book = self.download()
        dimensions = book.active.column_dimensions
        for index, column in enumerate(columns, start=1):
            self.assertGreaterEqual(dimensions[get_column_letter(index)].width, len(column) + 2)
        self.assertGreaterEqual(dimensions['B'].width, len(long_text) + 2)
        self.assertGreaterEqual(dimensions['AI'].width, len(long_text * 2) + 2)
        self.assertEqual(book.active['AI4'].value, long_text * 2)
        book.close()

    def test_empty_import_column_widths_fit_headers(self):
        self.record.items.all().delete()
        _, book = self.download()
        for index, column in enumerate(self.columns, start=1):
            self.assertGreaterEqual(book.active.column_dimensions[get_column_letter(index)].width, len(column) + 2)
        book.close()
