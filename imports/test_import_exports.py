from datetime import timedelta
from io import BytesIO
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth.models import Permission, User
from django.test import TestCase, override_settings, tag
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .export_jobs import export_next_import, queue_export
from imports.models import Import, ImportExportJob, ImportItem, ImportType


class ImportExportTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('export-reader')
        cls.user.user_permissions.add(Permission.objects.get(codename='view_import'))
        cls.record = Import.objects.create(import_type=ImportType.objects.get(code='payments'),
            file_name='payments.xlsx', status=Import.Status.FAILED, metadata={'columns': ['ДБЗ', 'Сумма']})
        ImportItem.objects.create(import_record=cls.record, row_number=2,
            data={'ДБЗ': '00001', 'Сумма': 123}, status=ImportItem.Status.PROCESSED)
        ImportItem.objects.create(import_record=cls.record, row_number=4,
            data={'ДБЗ': '=1+1', 'Сумма': '0.01'}, status=ImportItem.Status.FAILED,
            error_message='=HYPERLINK("test")')

    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        settings = override_settings(MEDIA_ROOT=self.directory.name, IMPORT_REPORT_ROOT=self.directory.name)
        settings.enable()
        self.addCleanup(settings.disable)
        self.client.force_login(self.user)

class ImportExportTests(ImportExportTestCase):
    def queue(self, **params):
        return self.client.get(reverse('imports:download', args=[self.record.pk]), params,
            HTTP_X_IMPORT_EXPORT='1')

    def test_request_only_queues_and_reuses_active_job(self):
        with patch('imports.files.build_import_workbook') as builder:
            first = self.queue(with_errors='1')
            second = self.queue(with_errors='1')
        builder.assert_not_called()
        self.assertEqual(first.status_code, 202)
        self.assertEqual(first['Cache-Control'], 'no-store')
        self.assertEqual(first.json(), second.json())
        self.assertEqual(ImportExportJob.objects.count(), 1)
        job = ImportExportJob.objects.get()
        self.assertTrue(job.include_errors)
        self.assertEqual(job.progress, 0)
        self.assertFalse(job.file)

    def test_worker_reports_progress_and_serves_saved_file_without_rebuilding(self):
        result = self.queue(with_errors='1').json()
        job = ImportExportJob.objects.get(pk=result['job_id'])
        observed = []
        from .files import build_import_workbook

        def build(*args, **kwargs):
            callback = kwargs['progress']

            def observe(percent, stage):
                callback(percent, stage)
                job.refresh_from_db()
                observed.append((job.progress, job.stage))

            kwargs['progress'] = observe
            return build_import_workbook(*args, **kwargs)

        with patch('imports.export_jobs.build_import_workbook', side_effect=build):
            self.assertTrue(export_next_import())
        job.refresh_from_db()
        self.assertEqual(job.status, ImportExportJob.Status.COMPLETED)
        self.assertEqual(job.progress, 100)
        self.assertTrue(job.report_file)
        self.assertIn((95, 'Упаковка XLSX'), observed)
        self.assertEqual([p for p, _ in observed], sorted(p for p, _ in observed))
        status = self.client.get(result['status_url'])
        self.assertEqual(status['Cache-Control'], 'no-store')
        with patch('imports.files.build_import_workbook') as builder:
            response = self.client.get(status.json()['download_url'])
            content = b''.join(response.streaming_content)
        builder.assert_not_called()
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(content))
        self.addCleanup(book.close)
        self.assertEqual(book.active['A2'].value, '00001')
        self.assertIsNone(book.active['A3'].value)
        self.assertEqual(book.active['A4'].data_type, 's')
        self.assertEqual(book.active['C4'].value, '=HYPERLINK("test")')
        self.assertEqual(book.active['C4'].data_type, 's')
        self.assertFalse(export_next_import())

    def test_jobs_and_files_are_not_accessible_to_another_user(self):
        result = self.queue().json()
        export_next_import()
        owner_status = self.client.get(result['status_url']).json()
        other = User.objects.create_user('other-reader')
        other.user_permissions.add(Permission.objects.get(codename='view_import'))
        self.client.force_login(other)
        self.assertEqual(self.client.get(result['status_url']).status_code, 404)
        self.assertEqual(self.client.get(owner_status['download_url']).status_code, 404)

    def test_not_ready_file_cannot_be_downloaded(self):
        job = queue_export(self.record.pk, user=self.user, include_errors=True)
        self.assertEqual(self.client.get(reverse('imports:export_file', args=[job.pk])).status_code, 404)

    def test_changing_import_cannot_be_exported(self):
        for status in (Import.Status.NEW, Import.Status.PROCESSING, Import.Status.IMPORTING):
            self.record.status = status
            self.record.save(update_fields=['status'])
            self.assertEqual(self.queue().status_code, 409)
        self.assertFalse(ImportExportJob.objects.exists())

    def test_worker_failure_is_persisted_and_can_be_retried(self):
        job = queue_export(self.record.pk, user=self.user, include_errors=True)
        with patch('imports.export_jobs.build_import_workbook', side_effect=OSError('disk full')):
            with self.assertLogs('imports.export_jobs', level='ERROR'):
                export_next_import()
        job.refresh_from_db()
        self.assertEqual(job.status, ImportExportJob.Status.FAILED)
        self.assertTrue(job.error_message)
        self.assertFalse(job.file)
        new_job = queue_export(self.record.pk, user=self.user, include_errors=True)
        self.assertNotEqual(job.pk, new_job.pk)

    def test_revoked_permission_prevents_background_export(self):
        job = queue_export(self.record.pk, user=self.user, include_errors=True)
        self.user.user_permissions.clear()
        with self.assertLogs('imports.export_jobs', level='ERROR'):
            export_next_import()
        job.refresh_from_db()
        self.assertEqual(job.status, ImportExportJob.Status.FAILED)
        self.assertFalse(job.file)

    def test_worker_recovers_stale_job_but_does_not_claim_running_job(self):
        job = queue_export(self.record.pk, user=self.user, include_errors=True)
        job.status = ImportExportJob.Status.PROCESSING
        job.token = 'old-token'
        job.heartbeat = timezone.now()
        job.save()
        self.assertFalse(export_next_import())
        job.heartbeat = timezone.now() - timedelta(minutes=11)
        job.save(update_fields=['heartbeat'])
        self.assertTrue(export_next_import())
        job.refresh_from_db()
        self.assertEqual(job.status, ImportExportJob.Status.COMPLETED)
        self.assertNotEqual(job.token, 'old-token')


@tag('slow')
class LargeImportExportTests(ImportExportTestCase):
    def test_100000_rows_exported_completely_with_bounded_width_sampling(self):
        self.record.items.all().delete()
        columns = ['ДБЗ', 'ИИН', 'ФИО', 'Сумма', 'Дата']
        self.record.metadata = {'columns': columns}
        self.record.save(update_fields=['metadata'])
        for start in range(0, 100000, 1000):
            ImportItem.objects.bulk_create([
                ImportItem(import_record=self.record, row_number=index + 2,
                    data={'ДБЗ': f'{index:012d}', 'ИИН': '000000000001',
                        'ФИО': f'Заёмщик {index}', 'Сумма': '123.45', 'Дата': '2026-10-08'},
                    status=ImportItem.Status.FAILED, error_message=f'Ошибка строки {index + 2}')
                for index in range(start, start + 1000)
            ], batch_size=1000)
        job = queue_export(self.record.pk, user=self.user, include_errors=True)
        from .files import build_import_workbook
        with patch('imports.export_jobs.build_import_workbook', wraps=build_import_workbook) as builder:
            self.assertTrue(export_next_import())
        self.assertEqual(builder.call_args.kwargs['width_sample'], 1000)
        job.refresh_from_db()
        self.assertEqual(job.status, ImportExportJob.Status.COMPLETED)
        self.assertEqual(job.progress, 100)
        with job.report_file.open('rb') as content:
            book = load_workbook(content, read_only=True)
            try:
                count = 0
                for row in book.active.iter_rows(values_only=True):
                    count += 1
                self.assertEqual(count, 100001)
                self.assertEqual(row[0], '000000099999')
                self.assertEqual(row[-1], 'Ошибка строки 100001')
            finally:
                book.close()
